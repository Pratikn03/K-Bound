import importlib.util
from pathlib import Path
import unittest
import tempfile
from io import BytesIO
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from zipfile import ZipFile
import hashlib

MODULE = Path(__file__).resolve().parents[1] / 'docs/research/kbound/scripts' / 'word_export.py'
spec = importlib.util.spec_from_file_location('word_export', MODULE) if MODULE.exists() else None
adapter = importlib.util.module_from_spec(spec) if spec else None
if spec:
    spec.loader.exec_module(adapter)

class ExportContracts(unittest.TestCase):
    def api(self):
        self.assertIsNotNone(adapter, 'staging adapter is not implemented')
        return adapter

    def test_main_source_selection_rejects_combined_driver(self):
        a = self.api()
        with self.assertRaises(ValueError):
            a.select_source(Path('/tmp/kbound_main_with_appendix.tex'), 'main')
        self.assertEqual(a.select_source(Path('/tmp/kbound_short_main.tex'), 'main').name, 'kbound_short_main.tex')

    def test_combined_source_selection_rejects_main_driver(self):
        a = self.api()
        with self.assertRaises(ValueError):
            a.select_source(Path('/tmp/kbound_short_main.tex'), 'combined')

    def test_source_tables_cannot_silently_disappear_in_parser(self):
        a = self.api()
        with self.assertRaisesRegex(ValueError, 'table'):
            a.source_contract(r'\begin{tabular}{c}x\end{tabular}', {'blocks':[]})

    def test_contract_counts_native_math_and_requires_three_figures(self):
        a = self.api()
        ast = {'blocks':[{'t':'Image'}, {'t':'Image'}, {'t':'Image'}, {'t':'Table'}, {'t':'Math'}, {'t':'Math'}]}
        self.assertEqual(a.source_contract(r'\begin{tabular}{c}x\end{tabular}', ast), {'tables':1,'figures':3,'math':2})
        with self.assertRaisesRegex(ValueError, 'figure'):
            a.source_contract('', {'blocks':[{'t':'Math'}]})

    def test_native_math_loss_is_hard_failure(self):
        a = self.api()
        with self.assertRaisesRegex(ValueError, 'math'):
            a.validate_counts({'tables':1,'figures':3,'math':1}, {'tables':1,'figures':3,'math':2})

    def test_table_loss_and_figure_loss_are_hard_failures(self):
        a = self.api()
        for key in ('tables','figures'):
            observed={'tables':1,'figures':3,'math':2}
            observed[key]-=1
            with self.assertRaisesRegex(ValueError,key):
                a.validate_counts(observed, {'tables':1,'figures':3,'math':2})

    def test_pandoc_warning_is_never_suppressed(self):
        a = self.api()
        with self.assertRaisesRegex(RuntimeError,'warning'):
            a.require_clean_process(0, 'ok', '[WARNING] math could not be converted')
        with self.assertRaises(RuntimeError):
            a.require_clean_process(1, '', '')
        self.assertEqual(a.require_clean_process(0,'ok',''), 'ok')

    def test_existing_output_is_refused(self):
        a = self.api()
        with self.assertRaises(FileExistsError):
            a.require_fresh_output(Path(__file__))

    def test_image_path_description_is_replaced_and_other_package_bytes_preserved(self):
        a=self.api()
        self.assertTrue(hasattr(a,'sanitize_image_descriptions'),'exporter lacks local-path metadata sanitizer')
        xml=b'<w:document><w:drawing><wp:docPr descr="Meaningful flow diagram"/><pic:cNvPr descr="/Volumes/T9/private/current-flow.png"/></w:drawing><w:t>Keep /Volumes/T9 as intentional prose</w:t><m:oMath>equation</m:oMath></w:document>'
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            original=Path(tmp)/'original.docx'
            output=Path(tmp)/'clean.docx'
            with ZipFile(original,'w') as z:
                z.writestr('word/document.xml',xml)
                z.writestr('word/media/image1.png',b'unchanged image bytes')
                z.writestr('word/_rels/document.xml.rels',b'https://example.com/science')
            before=original.read_bytes()
            self.assertEqual(a.sanitize_image_descriptions(original,output),1)
            self.assertEqual(original.read_bytes(),before)
            with ZipFile(output) as z:
                self.assertEqual(z.read('word/document.xml'),xml.replace(b'/Volumes/T9/private/current-flow.png',b'Meaningful flow diagram'))
                self.assertEqual(z.read('word/media/image1.png'),b'unchanged image bytes')
                self.assertEqual(z.read('word/_rels/document.xml.rels'),b'https://example.com/science')
            with self.assertRaises(FileExistsError):
                a.sanitize_image_descriptions(original,output)

    def test_image_url_and_meaningful_description_are_not_scrubbed(self):
        a=self.api()
        self.assertTrue(hasattr(a,'sanitize_image_descriptions'),'exporter lacks local-path metadata sanitizer')
        xml=b'<w:document><w:drawing><wp:docPr descr="Meaningful flow diagram"/><pic:cNvPr descr="https://example.com/plot.png"/></w:drawing><w:drawing><pic:cNvPr descr="Comparison across frozen and adapted states"/></w:drawing></w:document>'
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            original=Path(tmp)/'original.docx'
            output=Path(tmp)/'clean.docx'
            with ZipFile(original,'w') as z:
                z.writestr('word/document.xml',xml)
            self.assertEqual(a.sanitize_image_descriptions(original,output),0)
            with ZipFile(output) as z:
                self.assertEqual(z.read('word/document.xml'),xml)

    def test_stale_optional_word_statistics_removed_without_inventing_counts(self):
        a=self.api()
        app=b'<Properties><Application>Microsoft Word</Application><Pages>1</Pages><Words>83</Words><Characters>472</Characters><Lines>7</Lines><Paragraphs>1</Paragraphs><CharactersWithSpaces>553</CharactersWithSpaces><Company>Intentional organization</Company></Properties>'
        expected=b'<Properties><Application>Microsoft Word</Application><Company>Intentional organization</Company></Properties>'
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            original=Path(tmp)/'original.docx'
            output=Path(tmp)/'clean.docx'
            with ZipFile(original,'w') as z:
                z.writestr('docProps/app.xml',app)
            a.sanitize_image_descriptions(original,output)
            with ZipFile(output) as z:
                self.assertEqual(z.read('docProps/app.xml'),expected)

    def test_optional_caption_is_numbered_without_losing_long_text(self):
        a=self.api()
        text=r'\begin{table}\caption[Short]{Long {caption} with $x$}\label{tab:a}\end{table}'
        fixed=a.normalize_word_tables(text)
        numbered,labels=a.legacy.number_environments(fixed)
        self.assertIn(r'\caption{Table 1. Long {caption} with $x$}',numbered)
        self.assertEqual(labels['tab:a'],'1')

    def test_action_table_misdeclared_column_is_rejected_not_hidden(self):
        a=self.api()
        text=r'\begin{tabular}{@{}p{0.18\textwidth}rrrrrccp{0.26\textwidth}@{}}'+'\n'+r'a & b & c & d & e & f & g & h \\'+'\n'+r'\end{tabular}'
        with self.assertRaisesRegex(ValueError,'misdeclared'):
            a.normalize_word_tables(text)
        corrected=text.replace('rrrrrcc','rrrrcc')
        self.assertEqual(a.normalize_word_tables(corrected),corrected)

    def test_algorithm_numbering_has_explicit_gap_without_changing_text(self):
        a=self.api()
        doc=Document()
        doc.add_paragraph('Algorithm 1. Gate',style='Heading 2')
        p=doc.add_paragraph('Else if U < 0:')
        num=OxmlElement('w:numPr')
        p._p.get_or_add_pPr().append(num)
        doc.add_paragraph('Next',style='Heading 1')
        a.format_algorithm_numbering(doc)
        self.assertEqual(p.text,'\u00a0Else if U < 0:')
        self.assertIn('w:pos="720"',p._p.xml)
        self.assertIn('w:hanging="360"',p._p.xml)

    def test_main_companion_refs_do_not_require_absent_appendix_labels(self):
        a=self.api()
        text=r'\newcommand{\KBSuppRef}[3]{#3}'+'\n'+r'\providecommand{\KBSuppRef}[3]{#1~\ref{#2}}'+'\n'+r'See \KBSuppRef{Appendix}{app:absent}{companion {details}}.'
        self.assertEqual(a.role_references(text,'main').strip(),'See companion {details}.')
        combined=a.role_references(text,'combined')
        self.assertIn(r'Appendix~\ref{app:absent}',combined)

    def test_converted_algorithm_reference_retains_algorithm_number(self):
        a=self.api()
        text=r'See Algorithm~\ref{alg:test}. \begin{algorithm}\caption{Gate}\label{alg:test}\begin{algorithmic}\Require Input\State Act\end{algorithmic}\end{algorithm}'
        converted=a.legacy.replace_algorithm_for_word(text)
        self.assertIn('See Algorithm~1.',a.legacy.resolve_cross_references(converted))

    def test_layout_conditionals_select_one_balanced_float_branch(self):
        a=self.api()
        text=r'\ifdefined\KBoundTMLRLayout\begin{figure}[H]\else\begin{figure}[!t]\fi CONTENT\end{figure}'
        self.assertEqual(a.layout_conditionals(text),r'\begin{figure}[!t] CONTENT\end{figure}')

    def test_two_current_figures_use_explicit_contract_not_obsolete_third(self):
        a=self.api()
        ast={'blocks':[{'t':'Image'},{'t':'Image'},{'t':'Math'}]}
        self.assertEqual(a.source_contract('',ast,expected_figures=2),{'tables':0,'figures':2,'math':1})

    def test_tikz_replacement_retains_caption_and_labels(self):
        a=self.api()
        text=r'\begin{figure*}\resizebox{0.78\textwidth}{!}{\begin{tikzpicture}abc\end{tikzpicture}}\caption{Flow}\label{fig:flow}\end{figure*}'
        out=a.replace_tikz(text,Path('/tmp/current-flow.png'))
        self.assertIn(r'\includegraphics[width=0.98\textwidth]{/tmp/current-flow.png}',out)
        self.assertIn(r'\caption{Flow}\label{fig:flow}',out)
        self.assertNotIn('tikzpicture',out)

    def test_media_bytes_must_match_current_source_images(self):
        a=self.api()
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            path=Path(tmp)/'package.docx'
            with ZipFile(path,'w') as z:
                z.writestr('word/media/image1.png',b'current image')
            correct=hashlib.sha256(b'current image').hexdigest()
            self.assertEqual(a.validate_media(path,[correct]),[correct])
            with self.assertRaisesRegex(ValueError,'media'):
                a.validate_media(path,[hashlib.sha256(b'obsolete image').hexdigest()])

    def test_combined_main_equation_macros_resolve_to_real_labels(self):
        a=self.api()
        text=r'\providecommand{\KBMainRef}[3]{#1~\ref{#2}}'+'\n'+r'\providecommand{\KBMainEqRef}[2]{Eq.~\eqref{#1}}'+'\n'+r'\KBMainRef{Theorem}{thm:a}{fallback} and \KBMainEqRef{eq:a}{fallback}'
        self.assertEqual(a.role_references(text,'combined').strip(),r'Theorem~\ref{thm:a} and Eq.~\eqref{eq:a}')

    def test_explicit_twelve_table_role_passes_without_changing_legacy_floor(self):
        a=self.api()
        doc=Document()
        doc.add_paragraph('iWildCam withheld')
        for _ in range(12):
            doc.add_table(rows=1, cols=1).cell(0,0).text='value'
        for _ in range(3):
            buf=BytesIO()
            Image.new('RGB',(2,2),'white').save(buf,format='PNG')
            buf.seek(0)
            doc.add_picture(buf)
        doc.add_paragraph('References',style='Heading 1')
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            p=Path(tmp)/'fixture.docx'
            doc.save(p)
            with self.assertRaisesRegex(RuntimeError,'at least 13'):
                a.legacy.validate_docx(p,0)
            a.legacy.validate_docx(p,0,expected_tables=12)
            with self.assertRaisesRegex(RuntimeError,'expected 11'):
                a.legacy.validate_docx(p,0,expected_tables=11)

if __name__ == '__main__':
    unittest.main()
